from __future__ import annotations

import collections
import datetime as dt
import subprocess
from pathlib import Path

import openpyxl
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "outputs" / "B13_Test_Report_Official"
UAT_XLSX = Path(
    r"C:\Users\Pongo\Documents\ALL_TEL-U_SEMESTER 6\Pengujian PL\Kuesioner User Acceptance Testing (UAT) Sistem Gabara LMS (Jawaban).xlsx"
)
JOURNAL_DOCX = Path(
    r"C:\Users\Pongo\Documents\ALL_TEL-U_SEMESTER 6\Tata Tulis Ilmiah\Jurnal Gabara - DRAFTING SITASI.docx"
)

LIKERT = {
    "Sangat Setuju (SS)": 5,
    "Setuju (S)": 4,
    "Netral (N)": 3,
    "Tidak Setuju (TS)": 2,
    "Tidak Setuju (S)": 2,
    "Sangat Tidak Setuju (STS)": 1,
}

REFERENSI_STANDAR = [
    [
        "ISO/IEC/IEEE 29119-3:2021",
        "Acuan struktur dokumentasi pengujian perangkat lunak.",
        "https://www.iso.org/standard/79429.html",
    ],
    [
        "ISO/IEC 25010:2023",
        "Acuan karakteristik kualitas produk perangkat lunak.",
        "https://www.iso.org/standard/78176.html",
    ],
    [
        "ISO 9241-11:2018",
        "Acuan konsep usability: efektivitas, efisiensi, dan kepuasan pengguna.",
        "https://www.iso.org/standard/63500.html",
    ],
]


def cmd(args: list[str]) -> str:
    try:
        return subprocess.check_output(args, cwd=ROOT, text=True, stderr=subprocess.STDOUT).strip()
    except Exception:
        return "-"


def kategori(nilai: float) -> str:
    if nilai >= 4.21:
        return "Sangat Baik"
    if nilai >= 3.41:
        return "Baik"
    if nilai >= 2.61:
        return "Cukup"
    if nilai >= 1.81:
        return "Kurang"
    return "Sangat Kurang"


def normalisasi_role(role: str | None) -> str:
    teks = (role or "").strip().lower()
    if "admin" in teks:
        return "Admin"
    if "mentor" in teks:
        return "Mentor"
    return "Murid/Siswa"


def rata(nilai: list[float]) -> float:
    return sum(nilai) / len(nilai)


def muat_uat() -> dict:
    wb = openpyxl.load_workbook(UAT_XLSX, data_only=True)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    rows = [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True) if any(row)]
    role_col = "Peran dalam komunitas Gabara"
    durasi_col = "Berapa lama anda telah menggunakan LMS Gabara?"
    pertanyaan = header[6:24]
    aspek = {
        "Fungsionalitas Sistem": pertanyaan[0:5],
        "Usability / Kemudahan Penggunaan": pertanyaan[5:10],
        "Desain Antarmuka": pertanyaan[10:14],
        "Manfaat Sistem": pertanyaan[14:18],
    }
    for row in rows:
        row["_role"] = normalisasi_role(row.get(role_col))

    aspek_rows = []
    for nama, cols in aspek.items():
        vals = [
            LIKERT.get(row.get(col))
            for row in rows
            for col in cols
            if LIKERT.get(row.get(col)) is not None
        ]
        avg = rata(vals)
        aspek_rows.append(
            {
                "aspek": nama,
                "butir": len(cols),
                "rata": avg,
                "persen": avg / 5 * 100,
                "kategori": kategori(avg),
            }
        )

    role_rows = []
    for role in ["Admin", "Mentor", "Murid/Siswa"]:
        vals = [
            LIKERT.get(row.get(col))
            for row in rows
            if row["_role"] == role
            for col in pertanyaan
            if LIKERT.get(row.get(col)) is not None
        ]
        avg = rata(vals)
        role_rows.append(
            {
                "role": role,
                "jumlah": sum(1 for row in rows if row["_role"] == role),
                "rata": avg,
                "persen": avg / 5 * 100,
                "kategori": kategori(avg),
            }
        )

    pertanyaan_rows = []
    for q in pertanyaan:
        vals = [LIKERT.get(row.get(q)) for row in rows if LIKERT.get(row.get(q)) is not None]
        avg = rata(vals)
        pertanyaan_rows.append({"pertanyaan": q, "rata": avg, "persen": avg / 5 * 100})

    semua = [
        LIKERT.get(row.get(q))
        for row in rows
        for q in pertanyaan
        if LIKERT.get(row.get(q)) is not None
    ]
    total = rata([row["rata"] for row in aspek_rows])

    return {
        "jumlah_responden": len(rows),
        "jumlah_pertanyaan": len(pertanyaan),
        "jumlah_jawaban": len(semua),
        "role_mentah": collections.Counter(row.get(role_col) for row in rows),
        "durasi": collections.Counter(row.get(durasi_col) for row in rows),
        "aspek": aspek_rows,
        "role": role_rows,
        "pertanyaan": pertanyaan_rows,
        "total": total,
        "total_persen": total / 5 * 100,
        "total_kategori": kategori(total),
        "rata_jawaban": rata(semua),
    }


def buat_style() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "judul_sampul": ParagraphStyle(
            "JudulSampul",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=21,
            leading=26,
            textColor=colors.HexColor("#12355B"),
            alignment=TA_LEFT,
            spaceAfter=12,
        ),
        "subjudul": ParagraphStyle(
            "Subjudul",
            parent=base["Normal"],
            fontSize=10.5,
            leading=14,
            textColor=colors.HexColor("#4A5F73"),
            alignment=TA_LEFT,
            spaceAfter=16,
        ),
        "h1": ParagraphStyle(
            "H1",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=13.5,
            leading=16,
            textColor=colors.HexColor("#12355B"),
            spaceBefore=11,
            spaceAfter=6,
        ),
        "h2": ParagraphStyle(
            "H2",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13,
            textColor=colors.HexColor("#12355B"),
            spaceBefore=7,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8.6,
            leading=11.2,
            spaceAfter=5,
        ),
        "small": ParagraphStyle(
            "Small",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.2,
            leading=9.0,
        ),
        "center": ParagraphStyle(
            "Center",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.3,
            leading=9,
            alignment=TA_CENTER,
        ),
    }


S = buat_style()


def p(teks: object, style: str = "body") -> Paragraph:
    return Paragraph(str(teks), S[style])


def tabel(data: list[list[object]], widths=None, header: bool = True) -> Table:
    isi = [[v if hasattr(v, "wrap") else p(v, "small") for v in row] for row in data]
    t = Table(isi, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CAD5E2")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12355B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    t.setStyle(TableStyle(commands))
    return t


def bullets(items: list[str]) -> list[Paragraph]:
    return [p(f"- {item}") for item in items]


def section(judul: str, *content) -> list:
    return [p(judul, "h1"), *content]


def kontrol_dokumen(doc_id: str, judul: str, branch: str, commit: str) -> Table:
    return tabel(
        [
            ["Atribut", "Nilai"],
            ["ID Dokumen", doc_id],
            ["Judul Dokumen", judul],
            ["Proyek", "Gabara LMS - Garasi Belajar Banjarnegara"],
            ["Jenis Dokumen", "Laporan Hasil Pengujian Perangkat Lunak"],
            ["Versi", "1.0 Resmi"],
            ["Status", "Final untuk pengumpulan akademik"],
            ["Disusun untuk", "B13 Test Report - Pengujian Perangkat Lunak"],
            ["Branch / Commit", f"{branch} / {commit}"],
            ["Basis Data", "MySQL (dbs_gabara); basis data uji otomatis: dbs_gabara_testing"],
            ["Tanggal Penyusunan", dt.datetime.now().strftime("%d %B %Y")],
            ["Kerahasiaan", "Penggunaan akademik"],
        ],
        widths=[4.8 * cm, 11.4 * cm],
    )


def tabel_referensi() -> Table:
    rows = [["Referensi", "Pemanfaatan dalam Dokumen", "Sumber"]]
    rows.extend(REFERENSI_STANDAR)
    rows.extend(
        [
            [
                "Jurnal Gabara - DRAFTING SITASI.docx",
                "Acuan lokal untuk klasifikasi responden dan interpretasi UAT.",
                str(JOURNAL_DOCX),
            ],
            [
                "Kuesioner UAT Sistem Gabara LMS (Jawaban).xlsx",
                "Sumber utama data UAT dan jawaban skala Likert.",
                str(UAT_XLSX),
            ],
        ]
    )
    return tabel(rows, widths=[4.2 * cm, 7.0 * cm, 5.0 * cm])


def sampul(doc_id: str, judul: str, subjudul: str, branch: str, commit: str) -> list:
    return [
        Spacer(1, 1.0 * cm),
        p(judul, "judul_sampul"),
        p(subjudul, "subjudul"),
        tabel(
            [
                ["ID Dokumen", doc_id],
                ["Versi", "1.0 Resmi"],
                ["Proyek", "Gabara LMS"],
                ["Branch / Commit", f"{branch} / {commit}"],
                ["Tanggal", dt.datetime.now().strftime("%d %B %Y")],
            ],
            widths=[4.2 * cm, 10.5 * cm],
            header=False,
        ),
        Spacer(1, 0.7 * cm),
        p(
            "Dokumen ini disusun dalam bahasa Indonesia dengan struktur formal yang mengacu pada praktik dokumentasi "
            "pengujian perangkat lunak internasional, terutama ISO/IEC/IEEE 29119-3, serta disesuaikan dengan kebutuhan "
            "tugas B13 Test Report."
        ),
        PageBreak(),
    ]


def footer(judul: str):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#12355B"))
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawString(doc.leftMargin, A4[1] - 1.0 * cm, "Gabara LMS - Dokumentasi Pengujian Resmi")
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(A4[0] - doc.rightMargin, A4[1] - 1.0 * cm, judul[:70])
        canvas.drawString(doc.leftMargin, 0.82 * cm, f"Halaman {doc.page}")
        canvas.restoreState()

    return draw


def build(filename: str, judul: str, story: list) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / filename
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=1.65 * cm,
        rightMargin=1.65 * cm,
        topMargin=1.65 * cm,
        bottomMargin=1.35 * cm,
    )
    doc.build(story, onFirstPage=footer(judul), onLaterPages=footer(judul))
    return path


def laporan_uat(uat: dict, branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-UAT-LHP-001"
    judul = "Laporan Resmi Pengujian UAT"
    top = sorted(uat["pertanyaan"], key=lambda row: row["rata"], reverse=True)[:5]
    low = sorted(uat["pertanyaan"], key=lambda row: row["rata"])[:5]
    story = [
        *sampul(doc_id, judul, "Ringkasan penerimaan pengguna dengan klasifikasi tester sebagai Murid/Siswa.", branch, commit),
        *section("1. Kontrol Dokumen", kontrol_dokumen(doc_id, judul, branch, commit)),
        *section("2. Referensi Standar dan Referensi Proyek", tabel_referensi()),
        *section(
            "3. Ringkasan Eksekutif",
            p(
                f"Pengujian UAT melibatkan {uat['jumlah_responden']} responden dan {uat['jumlah_pertanyaan']} butir "
                f"pertanyaan tertutup. Responden dengan role mentah 'tester' dikategorikan sebagai Murid/Siswa karena "
                f"pengujian dilakukan dari sisi pengguna pembelajaran. Hasil akhir UAT adalah {uat['total']:.2f}/5,00 "
                f"atau {uat['total_persen']:.1f}% dan termasuk kategori {uat['total_kategori']}."
            ),
            tabel(
                [
                    ["Ukuran", "Hasil"],
                    ["Jumlah responden", "34 total: 5 Admin, 3 Mentor, 26 Murid/Siswa"],
                    ["Total jawaban Likert", uat["jumlah_jawaban"]],
                    ["Rata-rata keseluruhan UAT", f"{uat['total']:.2f}/5,00"],
                    ["Persentase kepuasan", f"{uat['total_persen']:.1f}%"],
                    ["Kategori", uat["total_kategori"]],
                    ["Aspek tertinggi", "Manfaat Sistem - 4,24/5,00"],
                    ["Aspek terendah", "Fungsionalitas Sistem - 4,00/5,00"],
                ],
                widths=[5.2 * cm, 11.0 * cm],
            ),
        ),
        *section(
            "4. Dasar dan Tujuan Pengujian",
            p(
                "Dasar pengujian mencakup implementasi LMS Gabara, kuesioner UAT, draf jurnal Gabara, dan kebutuhan "
                "fungsional sistem seperti manajemen kelas, materi, tugas, kuis, diskusi, dasbor, serta pemantauan "
                "peserta didik. Tujuan pengujian adalah menilai penerimaan pengguna terhadap LMS Gabara sebagai media "
                "pembelajaran non-formal."
            ),
        ),
        *section(
            "5. Ruang Lingkup",
            tabel(
                [
                    ["Termasuk", "Tidak Termasuk"],
                    [
                        "Persepsi kesesuaian fungsi, usability, desain antarmuka, manfaat sistem, analisis per role, dan masukan kualitatif.",
                        "Load testing formal, penetration testing, pengukuran hasil belajar jangka panjang, dan audit produksi.",
                    ],
                ],
                widths=[8.1 * cm, 8.1 * cm],
            ),
        ),
        *section(
            "6. Lingkungan dan Data Uji",
            tabel(
                [
                    ["Item", "Keterangan"],
                    ["Aplikasi", "Gabara LMS berbasis Laravel, React/Inertia, dan MySQL."],
                    ["Sumber data UAT", str(UAT_XLSX)],
                    ["Aturan klasifikasi", "Role mentah 'tester' dinormalisasi menjadi Murid/Siswa."],
                    ["Distribusi role", "Admin=5, Mentor=3, Murid/Siswa=26."],
                    ["Durasi penggunaan dominan", "Kurang dari 1 minggu: 23 responden."],
                ],
                widths=[4.6 * cm, 11.6 * cm],
            ),
        ),
        *section(
            "7. Pendekatan Pengujian",
            p(
                "Analisis UAT menggunakan pendekatan deskriptif kuantitatif. Jawaban skala Likert dikonversi menjadi "
                "skor 1 sampai 5. Rata-rata setiap aspek dihitung dari butir pertanyaan pada aspek tersebut. Rata-rata "
                "keseluruhan dihitung dari rata-rata empat aspek agar tiap dimensi memiliki bobot seimbang."
            ),
        ),
        *section(
            "8. Kriteria Masuk dan Kriteria Keluar",
            tabel(
                [
                    ["Jenis", "Kriteria", "Status"],
                    ["Masuk", "Data kuesioner UAT tersedia dan dapat dibaca.", "Terpenuhi"],
                    ["Masuk", "Aturan role tester sebagai Murid/Siswa disepakati.", "Terpenuhi"],
                    ["Keluar", "Hasil UAT minimal kategori Baik.", "Terpenuhi"],
                    ["Keluar", "Temuan dan rekomendasi terdokumentasi.", "Terpenuhi"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "9. Hasil Pengujian",
            p("9.1 Hasil per Aspek", "h2"),
            tabel(
                [["Aspek", "Butir", "Rata-rata", "Persentase", "Kategori"]]
                + [
                    [row["aspek"], row["butir"], f"{row['rata']:.2f}", f"{row['persen']:.1f}%", row["kategori"]]
                    for row in uat["aspek"]
                ],
                widths=[5.0 * cm, 1.5 * cm, 2.2 * cm, 2.5 * cm, 3.0 * cm],
            ),
            p("9.2 Hasil per Role", "h2"),
            tabel(
                [["Role", "Jumlah", "Rata-rata", "Persentase", "Kategori"]]
                + [
                    [row["role"], row["jumlah"], f"{row['rata']:.2f}", f"{row['persen']:.1f}%", row["kategori"]]
                    for row in uat["role"]
                ],
                widths=[3.2 * cm, 2.8 * cm, 2.4 * cm, 2.8 * cm, 3.0 * cm],
            ),
            p("9.3 Butir dengan Nilai Tertinggi", "h2"),
            tabel(
                [["Pertanyaan", "Rata-rata", "Persentase"]]
                + [[row["pertanyaan"], f"{row['rata']:.2f}", f"{row['persen']:.1f}%"] for row in top],
                widths=[11.2 * cm, 2.0 * cm, 2.4 * cm],
            ),
            p("9.4 Butir Prioritas Perbaikan", "h2"),
            tabel(
                [["Pertanyaan", "Rata-rata", "Persentase"]]
                + [[row["pertanyaan"], f"{row['rata']:.2f}", f"{row['persen']:.1f}%"] for row in low],
                widths=[11.2 * cm, 2.0 * cm, 2.4 * cm],
            ),
        ),
        *section(
            "10. Temuan, Risiko, dan Rekomendasi",
            tabel(
                [
                    ["ID", "Temuan", "Risiko", "Rekomendasi"],
                    [
                        "UAT-01",
                        "Skor Murid/Siswa lebih rendah dibanding Admin dan Mentor.",
                        "Hambatan pada alur belajar harian dapat menurunkan adopsi.",
                        "Prioritaskan enrollment, kuis, tugas, loading, dan navigasi siswa.",
                    ],
                    [
                        "UAT-02",
                        "Fungsionalitas menjadi aspek terendah.",
                        "Kendala fitur inti lebih terasa saat pembelajaran berlangsung.",
                        "Pertahankan pengujian regresi untuk kuis, tugas, dan enrollment.",
                    ],
                    [
                        "UAT-03",
                        "Masukan terbuka meminta video pembelajaran dan dark mode.",
                        "Ekspektasi pengguna dapat melampaui fitur saat ini.",
                        "Rencanakan pengembangan berdasarkan prioritas dampak.",
                    ],
                ],
                widths=[1.4 * cm, 5.0 * cm, 4.8 * cm, 5.0 * cm],
            ),
        ),
        *section(
            "11. Kesimpulan Resmi",
            p(
                "Hasil UAT mendukung penerimaan bersyarat LMS Gabara untuk penggunaan akademik dan operasional terbatas. "
                "Sistem diterima pada kategori Baik, dengan bukti terkuat pada aspek Manfaat Sistem. Perbaikan berikutnya "
                "sebaiknya berfokus pada pengalaman Murid/Siswa dan reliabilitas fitur inti."
            ),
        ),
    ]
    return build("01_Laporan_Resmi_UAT_Gabara.pdf", judul, story)


def laporan_fungsional(uat: dict, branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-FUNC-LHP-002"
    judul = "Laporan Resmi Pengujian Unit dan Pengujian Fungsional UI"
    story = [
        *sampul(doc_id, judul, "Laporan formal white-box unit test dan black-box feature/UI test.", branch, commit),
        *section("1. Kontrol Dokumen", kontrol_dokumen(doc_id, judul, branch, commit)),
        *section("2. Referensi Standar dan Referensi Proyek", tabel_referensi()),
        *section(
            "3. Ringkasan Eksekutif",
            p(
                "Paket pengujian fungsional mencakup unit test untuk logika internal dan feature/UI test untuk perilaku "
                "route aplikasi. Suite LegacyTestCaseAlignmentTest telah lulus dan menjadi bukti utama kesesuaian terhadap "
                "testcase lama. Full PHPUnit masih memerlukan pemeliharaan karena beberapa test historis belum sesuai dengan "
                "implementasi Inertia/controller saat ini."
            ),
        ),
        *section(
            "4. Item Uji",
            tabel(
                [
                    ["Item Uji", "Komponen / File", "Risiko yang Dicegah"],
                    ["Enrollment", "EnrollmentController, LegacyTestCaseAlignmentTest", "Kode tidak valid, kelas privat, duplicate enrollment, pesan sukses."],
                    ["Submission", "SubmissionController, LegacyTestCaseAlignmentTest", "Validasi file, deadline, update submission, batas nilai."],
                    ["Quiz", "QuizAttemptController, QuizController, QuizUnitTest", "Status draft, jadwal, batas attempt, scoring."],
                    ["Discussion", "DiscussionUnitTest", "Validasi enrollment, interval reply, status diskusi."],
                    ["Meeting", "MeetingUnitTest", "Regresi CRUD meeting, material, assignment."],
                    ["Auth/Profile", "Feature/Auth, ProfileTest", "Login, register, password, profile."],
                ],
                widths=[3.2 * cm, 5.6 * cm, 7.4 * cm],
            ),
        ),
        *section(
            "5. Pendekatan dan Lingkungan",
            tabel(
                [
                    ["Aspek", "Keterangan"],
                    ["Kotak putih / pengujian unit", "Memvalidasi aturan bisnis dan perhitungan skor dekat dengan implementasi."],
                    ["Kotak hitam / pengujian fitur", "Menguji route HTTP dan hasil sesi dari perspektif pengguna."],
                    ["Penjaga regresi", "LegacyTestCaseAlignmentTest menjaga perilaku testcase lama setelah perubahan branch."],
                    ["Objek tiruan pengujian", "Cloudinary upload API di-mock agar tidak bergantung jaringan eksternal."],
                    ["Database", "MySQL; phpunit.xml mengarah ke dbs_gabara_testing."],
                ],
                widths=[4.7 * cm, 11.5 * cm],
            ),
        ),
        *section(
            "6. Kriteria Masuk dan Kriteria Keluar",
            tabel(
                [
                    ["Jenis", "Kriteria", "Status"],
                    ["Masuk", "Dependency tersedia dan Laravel dapat bootstrap.", "Terpenuhi"],
                    ["Masuk", "Database MySQL testing tersedia.", "Terpenuhi"],
                    ["Keluar", "Legacy testcase alignment test lulus.", "Terpenuhi"],
                    ["Keluar", "Full PHPUnit suite lulus.", "Belum Terpenuhi"],
                    ["Keluar", "Temuan terdokumentasi.", "Terpenuhi"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "7. Hasil Eksekusi",
            tabel(
                [
                    ["Command", "Hasil", "Bukti"],
                    ["php artisan test --filter=LegacyTestCaseAlignmentTest", "Pass", "4 tests, 60 assertions."],
                    ["php artisan test", "Partial Fail", "30 passed, 13 failed pada area legacy yang terpisah."],
                    ["php artisan route:list --except-vendor", "Pass", "68 route terdeteksi."],
                    ["npm.cmd run build", "Pass", "4061 modules transformed; build selesai."],
                    ["npm.cmd run lint", "Fail", "ConfigError: plugin react-hooks tidak boleh didefinisikan ulang."],
                ],
                widths=[5.3 * cm, 2.4 * cm, 8.5 * cm],
            ),
        ),
        *section(
            "8. Traceability Ringkas",
            tabel(
                [
                    ["Fitur / Kebutuhan", "Bukti Otomatis", "Status"],
                    ["Enrollment", "LegacyTestCaseAlignmentTest", "Pass"],
                    ["Submission dan grading", "LegacyTestCaseAlignmentTest", "Pass"],
                    ["Quiz attempt dan skor", "LegacyTestCaseAlignmentTest + QuizUnitTest", "Pass"],
                    ["Discussion", "DiscussionUnitTest", "Pass"],
                    ["Meeting", "MeetingUnitTest", "Test perlu update"],
                    ["Auth/Profile", "Feature/Auth + ProfileTest", "Test perlu update"],
                ],
                widths=[5.2 * cm, 6.2 * cm, 4.0 * cm],
            ),
        ),
        *section(
            "9. Kesimpulan Resmi",
            p(
                "Bukti fungsional terfokus untuk alignment testcase lama diterima. Namun, full automated suite belum dapat "
                "dijadikan release gate bersih sebelum test historis dan konfigurasi lint diperbaiki."
            ),
        ),
    ]
    return build("02_Laporan_Resmi_Pengujian_Fungsional_UI_Gabara.pdf", judul, story)


def laporan_performance(branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-PERF-LHP-003"
    judul = "Laporan Resmi Pengujian Performa"
    story = [
        *sampul(doc_id, judul, "Laporan formal pengujian non-fungsional berbasis build production frontend.", branch, commit),
        *section("1. Kontrol Dokumen", kontrol_dokumen(doc_id, judul, branch, commit)),
        *section("2. Referensi Standar dan Referensi Proyek", tabel_referensi()),
        *section(
            "3. Ringkasan Eksekutif",
            p(
                "Pengujian performa dilakukan sebagai uji cepat build production. Aplikasi berhasil dibuild dalam 8,73 detik "
                "dan memproses 4061 modul frontend. Risiko utama berada pada chunk JavaScript ClassDetail yang besar dan "
                "beberapa aset gambar public/auth yang mendekati 800 kB."
            ),
        ),
        *section(
            "4. Tujuan Pengujian",
            *bullets(
                [
                    "Memastikan build production selesai tanpa error fatal TypeScript atau Vite.",
                    "Mengidentifikasi aset JavaScript, CSS, dan gambar yang berpotensi memperlambat load.",
                    "Memberikan rekomendasi optimasi sebelum deployment atau demo.",
                ]
            ),
        ),
        *section(
            "5. Ruang Lingkup",
            tabel(
                [
                    ["Termasuk", "Tidak Termasuk"],
                    [
                        "Status build, durasi build, jumlah modul, ukuran gzip, warning bundle, dan aset gambar besar.",
                        "Load test concurrent user, throttling jaringan, profiling server, profiling query database, dan audit Lighthouse.",
                    ],
                ],
                widths=[8.1 * cm, 8.1 * cm],
            ),
        ),
        *section(
            "6. Lingkungan dan Kriteria",
            tabel(
                [
                    ["Item", "Nilai"],
                    ["Command", "npm.cmd run build"],
                    ["Build tool", "Vite 7.1.9"],
                    ["Frontend stack", "React, Inertia, TypeScript"],
                    ["Kriteria keluar", "Build selesai; temuan aset besar terdokumentasi."],
                ],
                widths=[4.0 * cm, 12.2 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "7. Hasil Pengujian",
            tabel(
                [
                    ["Metrik", "Hasil", "Penilaian"],
                    ["Status", "Pass", "Build production selesai."],
                    ["Durasi build", "8,73 detik", "Layak untuk smoke test lokal."],
                    ["Modul diproses", "4061", "Dependency frontend cukup besar."],
                    ["CSS utama", "app-BgodYQzh.css 149,86 kB; gzip 22,57 kB", "Masih wajar."],
                    ["JS utama", "app-BdhcGsEf.js 366,05 kB; gzip 118,66 kB", "Di bawah threshold warning 500 kB."],
                    ["Chunk JS terbesar", "ClassDetail-AEVmJXht.js 610,38 kB; gzip 165,93 kB", "Melebihi threshold warning."],
                    ["Gambar besar", "image-login 795,79 kB; image-hero 787,36 kB; image-register 730,36 kB", "Perlu kompresi."],
                ],
                widths=[4.0 * cm, 6.4 * cm, 5.8 * cm],
            ),
        ),
        *section(
            "8. Temuan dan Rekomendasi",
            tabel(
                [
                    ["ID", "Severity", "Temuan", "Rekomendasi"],
                    ["PERF-01", "Medium", "Chunk ClassDetail melebihi 500 kB.", "Gunakan lazy loading route atau Rollup manualChunks."],
                    ["PERF-02", "Medium", "Gambar public/auth besar.", "Kompresi ke WebP/AVIF dan lazy-load bila memungkinkan."],
                    ["PERF-03", "Low", "Dependency graph memproses 4061 modul.", "Audit library yang tidak dipakai."],
                ],
                widths=[1.5 * cm, 2.0 * cm, 6.2 * cm, 6.5 * cm],
            ),
        ),
        *section(
            "9. Kesimpulan Resmi",
            p(
                "Pengujian cepat performa diterima secara bersyarat untuk kesiapan build. Project dapat dibuild, tetapi "
                "optimasi frontend tetap disarankan sebelum rilis atau demonstrasi yang sensitif terhadap performa."
            ),
        ),
    ]
    return build("03_Laporan_Resmi_Pengujian_Performa_Gabara.pdf", judul, story)


def laporan_regression(uat: dict, branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-REG-LHP-004"
    judul = "Laporan Resmi Pengujian Regresi"
    story = [
        *sampul(doc_id, judul, "Laporan formal penyelesaian pengujian regresi untuk stabilitas branch.", branch, commit),
        *section("1. Kontrol Dokumen", kontrol_dokumen(doc_id, judul, branch, commit)),
        *section("2. Referensi Standar dan Referensi Proyek", tabel_referensi()),
        *section(
            "3. Ringkasan Eksekutif",
            p(
                "Pengujian regresi menunjukkan bahwa suite penyelarasan testcase lama telah lulus setelah perubahan terakhir. "
                "Namun, release gate penuh belum bersih karena full PHPUnit dan konfigurasi ESLint masih gagal. Branch ini "
                "dapat dianggap stabil secara bersyarat hanya untuk cakupan testcase lama yang terdokumentasi."
            ),
        ),
        *section(
            "4. Ruang Lingkup dan Dasar Regresi",
            tabel(
                [
                    ["Area Regresi", "Dasar"],
                    ["Enrollment", "Pesan validasi dan aturan enrollment dari testcase lama."],
                    ["Assignment", "Validasi file, deadline, dan batas nilai."],
                    ["Quiz", "Status draft/publish, jadwal, limit attempt, scoring."],
                    ["Build", "Kompilasi frontend production."],
                    ["Route", "Inventaris route Laravel."],
                    ["Baseline UAT", f"Hasil UAT terkoreksi: {uat['total']:.2f}/5,00 ({uat['total_persen']:.1f}%)."],
                ],
                widths=[4.0 * cm, 12.2 * cm],
            ),
        ),
        *section(
            "5. Pendekatan Regresi",
            p(
                "Pendekatan regresi menggunakan pengujian backend otomatis untuk alur kerja murid yang berisiko tinggi, "
                "inventaris route untuk validasi permukaan aplikasi, build frontend sebagai uji cepat kompilasi, dan lint "
                "sebagai static quality gate. Checklist manual disediakan untuk alur yang belum sepenuhnya otomatis."
            ),
        ),
        *section(
            "6. Kriteria Masuk dan Kriteria Keluar",
            tabel(
                [
                    ["Jenis", "Kriteria", "Status"],
                    ["Masuk", "Branch dapat dibootstrap dan dependency tersedia.", "Terpenuhi"],
                    ["Masuk", "Database MySQL testing tersedia.", "Terpenuhi"],
                    ["Keluar", "Legacy testcase alignment test lulus.", "Terpenuhi"],
                    ["Keluar", "Frontend production build lulus.", "Terpenuhi"],
                    ["Keluar", "Full PHPUnit suite lulus.", "Belum Terpenuhi"],
                    ["Keluar", "Frontend lint lulus.", "Belum Terpenuhi"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "7. Hasil Gerbang Regresi",
            tabel(
                [
                    ["Gate", "Command / Bukti", "Status", "Disposition"],
                    ["Legacy testcase", "php artisan test --filter=LegacyTestCaseAlignmentTest", "Pass", "Diterima untuk alignment testcase lama."],
                    ["Full backend/UI otomatis", "php artisan test", "Partial Fail", "Perlu diperbaiki sebelum release gate."],
                    ["Route inventory", "php artisan route:list --except-vendor", "Pass", "68 route tersedia."],
                    ["Frontend build", "npm.cmd run build", "Pass", "Kesiapan build diterima."],
                    ["Frontend lint", "npm.cmd run lint", "Fail", "Konfigurasi ESLint perlu diperbaiki."],
                ],
                widths=[3.2 * cm, 5.4 * cm, 2.4 * cm, 5.2 * cm],
            ),
        ),
        *section(
            "8. Daftar Periksa Regresi Manual",
            tabel(
                [
                    ["Area", "Skenario", "Ekspektasi"],
                    ["Autentikasi", "Login, logout, registrasi, reset password.", "Validasi dan redirect berjalan benar."],
                    ["Peran", "Admin, mentor, dan siswa membuka dasbor.", "Navigasi dan proteksi sesuai role."],
                    ["Kelas", "Buat, ubah, hapus kelas dan enrollment.", "Data tersimpan dan akses benar."],
                    ["Meeting", "Create meeting, material, assignment.", "Konten terkait tampil di detail kelas."],
                    ["Assignment", "Upload/update submission dan grade.", "Pesan validasi dan batas nilai sesuai testcase."],
                    ["Quiz", "Start/continue/submit quiz.", "Jadwal, limit attempt, dan skor sesuai testcase."],
                    ["Discussion", "Buat thread, reply, tutup diskusi.", "Thread closed menolak reply baru."],
                ],
                widths=[2.6 * cm, 6.0 * cm, 7.6 * cm],
            ),
        ),
        *section(
            "9. Defect dan Risk Log",
            tabel(
                [
                    ["ID", "Severity", "Risiko / Defect", "Mitigasi"],
                    ["REG-01", "High", "Full PHPUnit gate gagal.", "Perbaiki stale tests dan missing ProfileController destroy behavior."],
                    ["REG-02", "Medium", "ESLint config gagal.", "Hapus duplicate react-hooks plugin declaration."],
                    ["REG-03", "Medium", "Chunk frontend besar.", "Split ClassDetail dan kompres gambar public/auth."],
                    ["REG-04", "Low", "Warning metadata doc-comment PHPUnit.", "Migrasi test ke attribute sebelum PHPUnit 12."],
                ],
                widths=[1.5 * cm, 2.0 * cm, 6.0 * cm, 6.7 * cm],
            ),
        ),
        *section(
            "10. Kesimpulan Resmi",
            p(
                "Pengujian regresi diterima hanya untuk cakupan penyelarasan testcase lama. Branch belum direkomendasikan "
                "sebagai kandidat rilis yang sepenuhnya bersih sampai full PHPUnit dan gerbang lint lulus."
            ),
        ),
    ]
    return build("04_Laporan_Resmi_Pengujian_Regresi_Gabara.pdf", judul, story)


def main() -> None:
    if not UAT_XLSX.exists():
        raise FileNotFoundError(UAT_XLSX)
    if not JOURNAL_DOCX.exists():
        raise FileNotFoundError(JOURNAL_DOCX)
    uat = muat_uat()
    branch = cmd(["git", "branch", "--show-current"])
    commit = cmd(["git", "rev-parse", "--short", "HEAD"])
    paths = [
        laporan_uat(uat, branch, commit),
        laporan_fungsional(uat, branch, commit),
        laporan_performance(branch, commit),
        laporan_regression(uat, branch, commit),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
