from __future__ import annotations

import collections
import datetime as dt
import subprocess
from pathlib import Path

import openpyxl
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "outputs" / "B13_Test_Report_Official"
UAT_XLSX = Path(
    r"C:\Users\Pongo\Documents\ALL_TEL-U_SEMESTER 6\Pengujian PL\Kuesioner User Acceptance Testing (UAT) Sistem Gabara LMS (Jawaban).xlsx"
)
JOURNAL_DOCX = Path(
    r"C:\Users\Pongo\Documents\ALL_TEL-U_SEMESTER 6\Tata Tulis Ilmiah\Jurnal Gabara - DRAFTING SITASI.docx"
)

LIKERT_SCORES = {
    "Sangat Setuju (SS)": 5,
    "Setuju (S)": 4,
    "Netral (N)": 3,
    "Tidak Setuju (TS)": 2,
    "Tidak Setuju (S)": 2,
    "Sangat Tidak Setuju (STS)": 1,
}

STANDARD_REFS = [
    [
        "ISO/IEC/IEEE 29119-3:2021",
        "Software and systems engineering - Software testing - Part 3: Test documentation.",
        "https://www.iso.org/standard/79429.html",
    ],
    [
        "ISO/IEC 25010:2023",
        "Systems and software Quality Requirements and Evaluation (SQuaRE) - Product quality model.",
        "https://www.iso.org/standard/78176.html",
    ],
    [
        "ISO 9241-11:2018",
        "Ergonomics of human-system interaction - Usability: Definitions and concepts.",
        "https://www.iso.org/standard/63500.html",
    ],
]


def command_value(args: list[str]) -> str:
    try:
        return subprocess.check_output(args, cwd=ROOT, text=True, stderr=subprocess.STDOUT).strip()
    except Exception:
        return "-"


def git_value(args: list[str]) -> str:
    return command_value(["git", *args])


def category(score: float) -> str:
    if score >= 4.21:
        return "Sangat Baik"
    if score >= 3.41:
        return "Baik"
    if score >= 2.61:
        return "Cukup"
    if score >= 1.81:
        return "Kurang"
    return "Sangat Kurang"


def normalize_role(role: str | None) -> str:
    text = (role or "").strip().lower()
    if "admin" in text:
        return "Admin"
    if "mentor" in text:
        return "Mentor"
    return "Murid"


def mean(values: list[float]) -> float:
    return sum(values) / len(values)


def load_uat() -> dict:
    wb = openpyxl.load_workbook(UAT_XLSX, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True) if any(row)]

    role_col = "Peran dalam komunitas Gabara"
    duration_col = "Berapa lama anda telah menggunakan LMS Gabara?"
    questions = headers[6:24]
    aspects = {
        "Fungsionalitas Sistem": questions[0:5],
        "Usability": questions[5:10],
        "Desain Antarmuka": questions[10:14],
        "Manfaat Sistem": questions[14:18],
    }

    for row in rows:
        row["_role"] = normalize_role(row.get(role_col))

    aspect_rows = []
    for aspect, cols in aspects.items():
        values = [
            LIKERT_SCORES.get(row.get(col))
            for row in rows
            for col in cols
            if LIKERT_SCORES.get(row.get(col)) is not None
        ]
        score = mean(values)
        aspect_rows.append(
            {
                "aspect": aspect,
                "items": len(cols),
                "score": score,
                "percent": score / 5 * 100,
                "category": category(score),
            }
        )

    role_rows = []
    for role in ["Admin", "Mentor", "Murid"]:
        values = [
            LIKERT_SCORES.get(row.get(col))
            for row in rows
            if row["_role"] == role
            for col in questions
            if LIKERT_SCORES.get(row.get(col)) is not None
        ]
        score = mean(values)
        role_rows.append(
            {
                "role": role,
                "count": sum(1 for row in rows if row["_role"] == role),
                "score": score,
                "percent": score / 5 * 100,
                "category": category(score),
            }
        )

    question_rows = []
    for question in questions:
        values = [
            LIKERT_SCORES.get(row.get(question))
            for row in rows
            if LIKERT_SCORES.get(row.get(question)) is not None
        ]
        score = mean(values)
        question_rows.append(
            {
                "question": question,
                "score": score,
                "percent": score / 5 * 100,
                "category": category(score),
            }
        )

    all_values = [
        LIKERT_SCORES.get(row.get(question))
        for row in rows
        for question in questions
        if LIKERT_SCORES.get(row.get(question)) is not None
    ]
    overall = mean([row["score"] for row in aspect_rows])
    raw_avg = mean(all_values)

    return {
        "respondent_count": len(rows),
        "question_count": len(questions),
        "answer_count": len(all_values),
        "roles_raw": collections.Counter(row.get(role_col) for row in rows),
        "roles": role_rows,
        "durations": collections.Counter(row.get(duration_col) for row in rows),
        "aspects": aspect_rows,
        "questions": question_rows,
        "overall": overall,
        "overall_percent": overall / 5 * 100,
        "overall_category": category(overall),
        "raw_avg": raw_avg,
        "raw_percent": raw_avg / 5 * 100,
    }


def styles():
    base = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle(
            "CoverTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=22,
            leading=27,
            textColor=colors.HexColor("#12355B"),
            alignment=TA_LEFT,
            spaceAfter=12,
        ),
        "cover_subtitle": ParagraphStyle(
            "CoverSubtitle",
            parent=base["Normal"],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#4A5F73"),
            alignment=TA_LEFT,
            spaceAfter=18,
        ),
        "h1": ParagraphStyle(
            "Heading1",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=17,
            textColor=colors.HexColor("#12355B"),
            spaceBefore=12,
            spaceAfter=6,
        ),
        "h2": ParagraphStyle(
            "Heading2",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#12355B"),
            spaceBefore=8,
            spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8.7,
            leading=11.3,
            alignment=TA_LEFT,
            spaceAfter=5,
        ),
        "small": ParagraphStyle(
            "Small",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.3,
            leading=9.1,
            alignment=TA_LEFT,
        ),
        "center": ParagraphStyle(
            "Center",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.3,
            alignment=TA_CENTER,
        ),
        "right": ParagraphStyle(
            "Right",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.3,
            alignment=TA_RIGHT,
        ),
    }


S = styles()


def par(text: object, style: str = "body") -> Paragraph:
    return Paragraph(str(text), S[style])


def cell(text: object, style: str = "small") -> Paragraph:
    return text if hasattr(text, "wrap") else par(text, style)


def table(data, widths=None, header=True, font_size=7.3) -> Table:
    converted = [[cell(value) for value in row] for row in data]
    t = Table(converted, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CAD5E2")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12355B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    t.setStyle(TableStyle(commands))
    return t


def bullets(items: list[str]) -> list[Paragraph]:
    return [par(f"- {item}") for item in items]


def section(title: str, *content) -> list:
    return [par(title, "h1"), *content]


def document_control(doc_id: str, title: str, branch: str, commit: str) -> Table:
    return table(
        [
            ["Field", "Value"],
            ["Document ID", doc_id],
            ["Document Title", title],
            ["Project", "Gabara LMS - Garasi Belajar Banjarnegara"],
            ["Document Type", "Software Test Report / Test Completion Report"],
            ["Version", "1.0 Official"],
            ["Status", "Final for academic submission"],
            ["Prepared For", "B13 Test Report - Pengujian Perangkat Lunak"],
            ["Branch / Commit", f"{branch} / {commit}"],
            ["Runtime Database", "MySQL (dbs_gabara); automated test database: dbs_gabara_testing"],
            ["Prepared Date", dt.datetime.now().strftime("%d %B %Y")],
            ["Confidentiality", "Academic use"],
        ],
        widths=[4.8 * cm, 11.4 * cm],
    )


def references_table(extra: list[list[str]] | None = None) -> Table:
    rows = [["Reference", "Use in This Document", "Source"]]
    for ref, use, source in STANDARD_REFS:
        rows.append([ref, use, source])
    rows.extend(
        [
            [
                "Jurnal Gabara - DRAFTING SITASI.docx",
                "Local research reference for UAT respondent classification and UAT interpretation.",
                str(JOURNAL_DOCX),
            ],
            [
                "Kuesioner User Acceptance Testing (UAT) Sistem Gabara LMS (Jawaban).xlsx",
                "Primary source for UAT respondent data and Likert responses.",
                str(UAT_XLSX),
            ],
        ]
    )
    if extra:
        rows.extend(extra)
    return table(rows, widths=[4.2 * cm, 7.0 * cm, 5.0 * cm])


def cover(doc_id: str, title: str, subtitle: str, branch: str, commit: str) -> list:
    return [
        Spacer(1, 1.0 * cm),
        par(title, "cover_title"),
        par(subtitle, "cover_subtitle"),
        table(
            [
                ["Document ID", doc_id],
                ["Version", "1.0 Official"],
                ["Project", "Gabara LMS"],
                ["Branch / Commit", f"{branch} / {commit}"],
                ["Prepared Date", dt.datetime.now().strftime("%d %B %Y")],
            ],
            widths=[4.2 * cm, 10.5 * cm],
            header=False,
        ),
        Spacer(1, 0.8 * cm),
        par(
            "This document follows a formal software testing report structure aligned with ISO/IEC/IEEE 29119-3 style "
            "documentation sections and maps quality interpretation to ISO/IEC 25010 where applicable."
        ),
        PageBreak(),
    ]


def footer(title: str):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#12355B"))
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawString(doc.leftMargin, A4[1] - 1.0 * cm, "Gabara LMS - Official Test Documentation")
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(A4[0] - doc.rightMargin, A4[1] - 1.0 * cm, title[:70])
        canvas.drawString(doc.leftMargin, 0.82 * cm, f"Page {doc.page}")
        canvas.restoreState()

    return draw


def build(filename: str, title: str, story: list) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / filename
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=1.65 * cm,
        rightMargin=1.65 * cm,
        topMargin=1.65 * cm,
        bottomMargin=1.35 * cm,
    )
    doc.build(story, onFirstPage=footer(title), onLaterPages=footer(title))
    return path


def acceptance_report(uat: dict, branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-UAT-TR-001"
    title = "Official Test Report - User Acceptance Testing (UAT)"
    top = sorted(uat["questions"], key=lambda row: row["score"], reverse=True)[:5]
    low = sorted(uat["questions"], key=lambda row: row["score"])[:5]
    roles = {row["role"]: row for row in uat["roles"]}

    story = [
        *cover(doc_id, title, "Formal acceptance test summary based on corrected student/respondent classification.", branch, commit),
        *section(
            "1. Document Control",
            document_control(doc_id, title, branch, commit),
        ),
        *section(
            "2. Normative and Project References",
            references_table(),
        ),
        *section(
            "3. Executive Summary",
            par(
                f"UAT was executed with {uat['respondent_count']} respondents and {uat['question_count']} closed "
                f"Likert-scale items. The corrected analysis classifies the tester respondent as Murid/Siswa because "
                f"the tester evaluated the system from the learner-side workflow. The final UAT score is "
                f"{uat['overall']:.2f}/5.00 ({uat['overall_percent']:.1f}%), categorized as {uat['overall_category']}."
            ),
            table(
                [
                    ["Measure", "Result"],
                    ["Respondents", "34 total: 5 Admin, 3 Mentor, 26 Murid/Siswa"],
                    ["Likert responses", f"{uat['answer_count']} scored responses"],
                    ["Overall UAT score", f"{uat['overall']:.2f}/5.00"],
                    ["Satisfaction percentage", f"{uat['overall_percent']:.1f}%"],
                    ["Result category", uat["overall_category"]],
                    ["Highest aspect", "Manfaat Sistem - 4.24/5.00"],
                    ["Lowest aspect", "Fungsionalitas Sistem - 4.00/5.00"],
                ],
                widths=[5.2 * cm, 11.0 * cm],
            ),
        ),
        *section(
            "4. Test Basis and Objectives",
            par(
                "The test basis consists of the Gabara LMS implementation, UAT questionnaire, project journal draft, "
                "and functional requirements for class management, material distribution, assignments, quizzes, "
                "discussion forum, dashboard, and learner monitoring. The objective is to verify stakeholder acceptance "
                "of Gabara LMS for non-formal learning activities at Garasi Belajar Banjarnegara."
            ),
        ),
        *section(
            "5. Scope",
            table(
                [
                    ["Included", "Excluded"],
                    [
                        "Functional suitability perception, usability, interface design, system benefit, role-based UAT analysis, qualitative feedback.",
                        "Formal load testing, security penetration testing, long-term learning outcome measurement, production incident analysis.",
                    ],
                ],
                widths=[8.1 * cm, 8.1 * cm],
            ),
        ),
        *section(
            "6. Test Environment and Data",
            table(
                [
                    ["Item", "Description"],
                    ["Application", "Gabara LMS web application using Laravel, React/Inertia, and MySQL."],
                    ["UAT data source", str(UAT_XLSX)],
                    ["Respondent classification", "Raw role 'tester' normalized into Murid/Siswa."],
                    [
                        "Role distribution",
                        f"Admin={roles['Admin']['count']}, Mentor={roles['Mentor']['count']}, Murid={roles['Murid']['count']}.",
                    ],
                    ["Dominant usage duration", "Kurang dari 1 minggu: 23 respondents."],
                ],
                widths=[4.6 * cm, 11.6 * cm],
            ),
        ),
        *section(
            "7. Test Approach",
            par(
                "The UAT uses descriptive quantitative analysis. Each Likert answer is converted to a score from 1 to 5. "
                "Question averages are grouped into four aspects. The overall score is calculated as the average of the "
                "four aspect scores so that each quality dimension has equal weight."
            ),
            table(
                [["Likert Option", "Score"], ["SS", "5"], ["S", "4"], ["N", "3"], ["TS", "2"], ["STS", "1"]],
                widths=[5.0 * cm, 2.0 * cm],
            ),
        ),
        *section(
            "8. Entry and Exit Criteria",
            table(
                [
                    ["Criterion Type", "Criterion", "Status"],
                    ["Entry", "UAT questionnaire responses available and readable.", "Met"],
                    ["Entry", "Respondent role mapping rule agreed: tester counted as Murid/Siswa.", "Met"],
                    ["Exit", "Overall UAT score is at least category Baik.", "Met"],
                    ["Exit", "Findings and improvement recommendations documented.", "Met"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "9. Test Results",
            par("9.1 Result by Aspect", "h2"),
            table(
                [["Aspect", "Items", "Score", "Percentage", "Category"]]
                + [
                    [row["aspect"], row["items"], f"{row['score']:.2f}", f"{row['percent']:.1f}%", row["category"]]
                    for row in uat["aspects"]
                ],
                widths=[5.0 * cm, 1.5 * cm, 2.2 * cm, 2.5 * cm, 3.0 * cm],
            ),
            par("9.2 Result by Role", "h2"),
            table(
                [["Role", "Respondents", "Score", "Percentage", "Category"]]
                + [
                    [row["role"], row["count"], f"{row['score']:.2f}", f"{row['percent']:.1f}%", row["category"]]
                    for row in uat["roles"]
                ],
                widths=[3.2 * cm, 2.8 * cm, 2.4 * cm, 2.8 * cm, 3.0 * cm],
            ),
            par("9.3 Highest Scoring Items", "h2"),
            table(
                [["Question", "Score", "Percentage"]]
                + [[row["question"], f"{row['score']:.2f}", f"{row['percent']:.1f}%"] for row in top],
                widths=[11.2 * cm, 2.0 * cm, 2.4 * cm],
            ),
            par("9.4 Priority Improvement Items", "h2"),
            table(
                [["Question", "Score", "Percentage"]]
                + [[row["question"], f"{row['score']:.2f}", f"{row['percent']:.1f}%"] for row in low],
                widths=[11.2 * cm, 2.0 * cm, 2.4 * cm],
            ),
        ),
        *section(
            "10. Findings, Risks, and Recommendations",
            table(
                [
                    ["ID", "Finding", "Risk", "Recommended Action"],
                    [
                        "UAT-01",
                        "Murid/Siswa score is lower than Admin and Mentor.",
                        "Learner-side friction may reduce daily adoption.",
                        "Prioritize learner workflows: enrollment, quiz, assignment, loading, and navigation.",
                    ],
                    [
                        "UAT-02",
                        "Functionality is the lowest aspect.",
                        "Core feature defects are more visible during actual learning.",
                        "Use legacy testcase regression tests for quiz, assignment, and enrollment.",
                    ],
                    [
                        "UAT-03",
                        "Feedback requests video learning and dark mode.",
                        "Future expectations may exceed current feature set.",
                        "Plan enhancements by release priority and impact.",
                    ],
                ],
                widths=[1.4 * cm, 5.0 * cm, 4.8 * cm, 5.0 * cm],
            ),
        ),
        *section(
            "11. Formal Conclusion",
            par(
                "The UAT result supports conditional acceptance of Gabara LMS for academic and limited operational use. "
                "The product is accepted in category Baik with the strongest evidence on system benefit. Improvements "
                "should focus on learner-side usability and functional reliability."
            ),
        ),
    ]
    return build("01_Official_UAT_Test_Report_Gabara.pdf", title, story)


def functionality_report(uat: dict, branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-FUNC-TR-002"
    title = "Official Test Report - Unit and UI Functionality Testing"
    story = [
        *cover(doc_id, title, "Formal functional test report covering white-box unit tests and black-box UI feature tests.", branch, commit),
        *section("1. Document Control", document_control(doc_id, title, branch, commit)),
        *section("2. Normative and Project References", references_table()),
        *section(
            "3. Executive Summary",
            par(
                "The formal functional test package contains unit tests for internal logic and feature/UI tests for route-level "
                "behavior. The legacy testcase alignment suite passes and provides focused evidence for the old testcase scope. "
                "The full PHPUnit suite still requires maintenance because several legacy tests no longer match the current "
                "Inertia/controller implementation."
            ),
        ),
        *section(
            "4. Test Items",
            table(
                [
                    ["Test Item", "Component / File", "Primary Risk Addressed"],
                    ["Enrollment", "EnrollmentController, LegacyTestCaseAlignmentTest", "Invalid code, private class, duplicate enrollment, success message."],
                    ["Submission", "SubmissionController, LegacyTestCaseAlignmentTest", "File validation, deadline, update, grade min/max."],
                    ["Quiz", "QuizAttemptController, QuizController, QuizUnitTest", "Draft visibility, schedule, max attempt, scoring."],
                    ["Discussion", "DiscussionUnitTest", "Enrollment validation, reply interval, closed discussion."],
                    ["Meeting", "MeetingUnitTest", "Meeting/material/assignment CRUD regression."],
                    ["Auth/Profile", "Feature/Auth, ProfileTest", "Login, register, password, profile workflows."],
                ],
                widths=[3.2 * cm, 5.6 * cm, 7.4 * cm],
            ),
        ),
        *section(
            "5. Scope",
            table(
                [
                    ["Included", "Excluded"],
                    [
                        "White-box unit logic, black-box route/feature behavior, validation messages, role-protected workflows, legacy testcase alignment.",
                        "Browser-driven end-to-end automation, visual regression screenshots, security testing, real external Cloudinary upload.",
                    ],
                ],
                widths=[8.1 * cm, 8.1 * cm],
            ),
        ),
        *section(
            "6. Test Approach",
            table(
                [
                    ["Technique", "Application"],
                    ["White-box / unit testing", "Directly validates business rules and scoring logic close to implementation."],
                    ["Black-box / feature testing", "Exercises HTTP routes and session outcomes from the user's perspective."],
                    ["Regression guard", "LegacyTestCaseAlignmentTest preserves old testcase behavior after branch changes."],
                    ["Test doubles", "Cloudinary upload API mocked to avoid external network dependency."],
                ],
                widths=[4.8 * cm, 11.4 * cm],
            ),
        ),
        *section(
            "7. Environment",
            table(
                [
                    ["Environment Item", "Value"],
                    ["Backend framework", "Laravel 12 / PHPUnit 11.5 via php artisan test."],
                    ["Frontend", "React 19, Inertia, Vite 7."],
                    ["Database", "MySQL; phpunit.xml points to dbs_gabara_testing."],
                    ["Routes", "68 application routes listed by php artisan route:list --except-vendor."],
                ],
                widths=[4.5 * cm, 11.7 * cm],
            ),
        ),
        *section(
            "8. Entry and Exit Criteria",
            table(
                [
                    ["Criterion Type", "Criterion", "Status"],
                    ["Entry", "Dependencies installed and Laravel app can bootstrap.", "Met"],
                    ["Entry", "MySQL testing database available.", "Met"],
                    ["Exit", "Legacy testcase alignment suite passes.", "Met"],
                    ["Exit", "Full PHPUnit suite passes.", "Not Met"],
                    ["Exit", "Known failures documented as defects/risks.", "Met"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "9. Test Execution Results",
            table(
                [
                    ["Command", "Result", "Evidence"],
                    ["php artisan test --filter=LegacyTestCaseAlignmentTest", "Pass", "4 tests, 60 assertions."],
                    ["php artisan test", "Partial Fail", "30 passed, 13 failed in unrelated legacy areas."],
                    ["php artisan route:list --except-vendor", "Pass", "68 routes listed."],
                    ["npm.cmd run build", "Pass", "4061 modules transformed; build completed."],
                    ["npm.cmd run lint", "Fail", "ConfigError: plugin react-hooks cannot be redefined."],
                ],
                widths=[5.3 * cm, 2.4 * cm, 8.5 * cm],
            ),
        ),
        *section(
            "10. Defect and Risk Log",
            table(
                [
                    ["ID", "Severity", "Observation", "Recommended Action"],
                    ["FUNC-01", "High", "Full PHPUnit suite is not clean.", "Repair MeetingUnitTest, Auth/Profile page expectations, and ProfileController delete route handling."],
                    ["FUNC-02", "Medium", "ESLint configuration fails before rule evaluation.", "Remove duplicate react-hooks plugin definition."],
                    ["FUNC-03", "Low", "PHPUnit doc-comment metadata warnings.", "Migrate @test doc-comments to PHPUnit attributes."],
                ],
                widths=[1.5 * cm, 2.0 * cm, 6.2 * cm, 6.5 * cm],
            ),
        ),
        *section(
            "11. Traceability Summary",
            table(
                [
                    ["Requirement / Feature", "Automated Evidence", "Status"],
                    ["Enrollment", "LegacyTestCaseAlignmentTest", "Pass"],
                    ["Assignment submission and grading", "LegacyTestCaseAlignmentTest", "Pass"],
                    ["Quiz attempt and score", "LegacyTestCaseAlignmentTest + QuizUnitTest", "Pass"],
                    ["Discussion", "DiscussionUnitTest", "Pass"],
                    ["Meeting", "MeetingUnitTest", "Failing tests need update"],
                    ["Auth/Profile", "Feature/Auth + ProfileTest", "Failing tests need update"],
                ],
                widths=[5.2 * cm, 6.2 * cm, 4.0 * cm],
            ),
        ),
        *section(
            "12. Formal Conclusion",
            par(
                "The focused functionality evidence for legacy testcase alignment is accepted. The broader automated suite "
                "is not yet release-ready because historical tests and configuration issues remain. These defects should be "
                "resolved before using the full suite as a release gate."
            ),
        ),
    ]
    return build("02_Official_Functionality_Test_Report_Gabara.pdf", title, story)


def performance_report(branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-PERF-TR-003"
    title = "Official Test Report - Performance Testing"
    story = [
        *cover(doc_id, title, "Formal non-functional smoke report for build performance and frontend asset size.", branch, commit),
        *section("1. Document Control", document_control(doc_id, title, branch, commit)),
        *section("2. Normative and Project References", references_table()),
        *section(
            "3. Executive Summary",
            par(
                "The performance test was conducted as a production-build smoke test. The application built successfully in "
                "8.73 seconds and processed 4061 frontend modules. The main performance risk is a large ClassDetail JavaScript "
                "chunk and several public/auth images approaching 800 kB."
            ),
        ),
        *section(
            "4. Test Objectives",
            *bullets(
                [
                    "Verify that the production frontend build completes without fatal TypeScript or Vite errors.",
                    "Identify large JavaScript/CSS/image assets that can affect load time.",
                    "Provide optimization recommendations before deployment or demo recording.",
                ]
            ),
        ),
        *section(
            "5. Scope",
            table(
                [
                    ["Included", "Excluded"],
                    [
                        "Build completion, build duration, module count, gzip sizes, bundle warning review, image asset review.",
                        "Concurrent user load test, network throttling, server response profiling, database query profiling, Lighthouse audit.",
                    ],
                ],
                widths=[8.1 * cm, 8.1 * cm],
            ),
        ),
        *section(
            "6. Test Environment",
            table(
                [
                    ["Item", "Value"],
                    ["Command", "npm.cmd run build"],
                    ["Build tool", "Vite 7.1.9"],
                    ["Frontend stack", "React, Inertia, TypeScript"],
                    ["Execution location", "Local project root on current branch."],
                ],
                widths=[4.0 * cm, 12.2 * cm],
            ),
        ),
        *section(
            "7. Entry and Exit Criteria",
            table(
                [
                    ["Criterion Type", "Criterion", "Status"],
                    ["Entry", "Node dependencies available.", "Met"],
                    ["Entry", "TypeScript and Vite configuration available.", "Met"],
                    ["Exit", "Production build completes successfully.", "Met"],
                    ["Exit", "High-risk asset findings documented.", "Met"],
                    ["Exit", "No bundle exceeds warning threshold.", "Not Met"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "8. Test Results",
            table(
                [
                    ["Metric", "Observed Result", "Assessment"],
                    ["Status", "Pass", "Production build completed."],
                    ["Build duration", "8.73 seconds", "Acceptable for local smoke test."],
                    ["Modules transformed", "4061", "Moderate dependency size."],
                    ["Main CSS", "app-BgodYQzh.css 149.86 kB; gzip 22.57 kB", "Acceptable."],
                    ["Main JS", "app-BdhcGsEf.js 366.05 kB; gzip 118.66 kB", "Below 500 kB warning threshold."],
                    ["Largest JS chunk", "ClassDetail-AEVmJXht.js 610.38 kB; gzip 165.93 kB", "Warning threshold exceeded."],
                    ["Large images", "image-login 795.79 kB; image-hero 787.36 kB; image-register 730.36 kB", "Optimize before deployment."],
                ],
                widths=[4.0 * cm, 6.4 * cm, 5.8 * cm],
            ),
        ),
        *section(
            "9. Findings and Recommendations",
            table(
                [
                    ["ID", "Severity", "Finding", "Recommendation"],
                    ["PERF-01", "Medium", "ClassDetail chunk exceeds 500 kB warning threshold.", "Apply route-level lazy loading or Rollup manualChunks."],
                    ["PERF-02", "Medium", "Public/auth images are large.", "Compress images to WebP/AVIF and lazy-load where possible."],
                    ["PERF-03", "Low", "Dependency graph includes 4061 transformed modules.", "Audit unused libraries and duplicate UI dependencies."],
                ],
                widths=[1.5 * cm, 2.0 * cm, 6.2 * cm, 6.5 * cm],
            ),
        ),
        *section(
            "10. Formal Conclusion",
            par(
                "Performance smoke testing is conditionally accepted for build readiness. The project is buildable, but "
                "frontend optimization should be completed before production release or performance-sensitive demonstrations."
            ),
        ),
    ]
    return build("03_Official_Performance_Test_Report_Gabara.pdf", title, story)


def regression_report(uat: dict, branch: str, commit: str) -> Path:
    doc_id = "GABARA-B13-REG-TR-004"
    title = "Official Test Report - Regression Testing"
    story = [
        *cover(doc_id, title, "Formal regression test completion report for branch stability and release gate evidence.", branch, commit),
        *section("1. Document Control", document_control(doc_id, title, branch, commit)),
        *section("2. Normative and Project References", references_table()),
        *section(
            "3. Executive Summary",
            par(
                "Regression testing confirms that the focused legacy testcase alignment suite passes after recent changes. "
                "However, the full release gate is not yet clean because historical unit/feature tests and ESLint configuration "
                "still fail. The branch should be treated as conditionally stable for the documented testcase scope only."
            ),
        ),
        *section(
            "4. Regression Scope and Test Basis",
            table(
                [
                    ["Regression Area", "Basis"],
                    ["Enrollment", "Legacy testcase messages and enrollment rules."],
                    ["Assignment", "File validation, deadline handling, grade validation."],
                    ["Quiz", "Draft/published status, schedule, attempt limit, scoring."],
                    ["Build", "Production frontend compilation."],
                    ["Routes", "Laravel route inventory."],
                    ["UAT baseline", f"Corrected UAT result: {uat['overall']:.2f}/5.00 ({uat['overall_percent']:.1f}%)."],
                ],
                widths=[4.0 * cm, 12.2 * cm],
            ),
        ),
        *section(
            "5. Test Approach",
            par(
                "The regression approach uses automated backend tests for high-risk learner workflows, route inventory for "
                "application surface validation, frontend build as a compilation smoke test, and lint as a static quality gate. "
                "Manual regression checklist items are provided for flows not yet fully automated."
            ),
        ),
        *section(
            "6. Entry and Exit Criteria",
            table(
                [
                    ["Criterion Type", "Criterion", "Status"],
                    ["Entry", "Branch can be bootstrapped and dependencies are installed.", "Met"],
                    ["Entry", "MySQL testing database is available.", "Met"],
                    ["Exit", "Legacy testcase alignment tests pass.", "Met"],
                    ["Exit", "Frontend production build passes.", "Met"],
                    ["Exit", "Full PHPUnit suite passes.", "Not Met"],
                    ["Exit", "Frontend lint passes.", "Not Met"],
                ],
                widths=[3.0 * cm, 10.0 * cm, 3.0 * cm],
            ),
        ),
        PageBreak(),
        *section(
            "7. Regression Gate Results",
            table(
                [
                    ["Gate", "Command / Evidence", "Status", "Disposition"],
                    ["Legacy testcase", "php artisan test --filter=LegacyTestCaseAlignmentTest", "Pass", "Accepted for old testcase alignment."],
                    ["Full automated backend/UI", "php artisan test", "Partial Fail", "Requires repair before release gate."],
                    ["Route inventory", "php artisan route:list --except-vendor", "Pass", "68 routes available."],
                    ["Frontend build", "npm.cmd run build", "Pass", "Build readiness accepted."],
                    ["Frontend lint", "npm.cmd run lint", "Fail", "ESLint config repair required."],
                ],
                widths=[3.2 * cm, 5.4 * cm, 2.4 * cm, 5.2 * cm],
            ),
        ),
        *section(
            "8. Manual Regression Checklist",
            table(
                [
                    ["Area", "Scenario", "Expected Result"],
                    ["Auth", "Login, logout, register, reset password.", "Validation and redirects work correctly."],
                    ["Role", "Admin, mentor, and student open dashboard.", "Role-specific navigation and protection work."],
                    ["Class", "Create/edit/delete class and enroll by code.", "Data persists and access rules hold."],
                    ["Meeting", "Create meeting, material, assignment.", "Related content appears in class detail."],
                    ["Assignment", "Upload/update submission and grade.", "Validation messages and grade boundaries match testcase."],
                    ["Quiz", "Start/continue/submit quiz.", "Schedule, attempt limit, and score behavior match testcase."],
                    ["Discussion", "Create thread, reply, close discussion.", "Closed threads reject new replies."],
                ],
                widths=[2.6 * cm, 6.0 * cm, 7.6 * cm],
            ),
        ),
        *section(
            "9. Defect and Risk Log",
            table(
                [
                    ["ID", "Severity", "Risk / Defect", "Mitigation"],
                    ["REG-01", "High", "Full PHPUnit gate fails.", "Repair stale tests and missing ProfileController destroy behavior or update route expectation."],
                    ["REG-02", "Medium", "ESLint config failure blocks static analysis.", "Remove duplicate react-hooks plugin declaration."],
                    ["REG-03", "Medium", "Large frontend chunk remains.", "Split ClassDetail and compress public/auth images."],
                    ["REG-04", "Low", "Warning from PHPUnit doc-comment metadata.", "Migrate tests to attributes before PHPUnit 12."],
                ],
                widths=[1.5 * cm, 2.0 * cm, 6.0 * cm, 6.7 * cm],
            ),
        ),
        *section(
            "10. Formal Conclusion",
            par(
                "Regression testing is accepted only for the focused legacy testcase alignment scope. The branch is not yet "
                "recommended as a fully clean release candidate until the full PHPUnit and lint gates pass."
            ),
        ),
    ]
    return build("04_Official_Regression_Test_Report_Gabara.pdf", title, story)


def main() -> None:
    if not UAT_XLSX.exists():
        raise FileNotFoundError(UAT_XLSX)
    if not JOURNAL_DOCX.exists():
        raise FileNotFoundError(JOURNAL_DOCX)

    uat = load_uat()
    branch = git_value(["branch", "--show-current"])
    commit = git_value(["rev-parse", "--short", "HEAD"])
    paths = [
        acceptance_report(uat, branch, commit),
        functionality_report(uat, branch, commit),
        performance_report(branch, commit),
        regression_report(uat, branch, commit),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
